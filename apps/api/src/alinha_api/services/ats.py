"""Deterministic ATS readability checks; no AI calls and no 'ATS approval' score.

Rules are common document-parsing precautions, not a reproduction of any
vendor's private screening algorithm. Never store the candidate's document.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Literal

from docx import Document
from pydantic import BaseModel, Field
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from alinha_api.services.document_reader import (
    MAX_UPLOAD_BYTES,
    extract_resume,
)

Kind = Literal["pdf", "docx", "txt", "text"]
Severity = Literal["important", "attention", "suggestion"]

EMAIL_RE = re.compile(r"(?<![\w.+-])[\w.+-]+@[\w.-]+\.[a-zA-Z]{2,}(?![\w.-])")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[ .-]?)?(?:\(?\d{2,3}\)?[ .-]?)?\d{4,5}[ .-]?\d{4}")
SECTION_PATTERNS: dict[str, re.Pattern[str]] = {
    "experience": re.compile(r"(?im)^\s*(?:experiências? profiss(?:ional|ionais)|histórico profissional|professional experience|work experience|empregos|carreira)\s*:?\s*$"),
    "education": re.compile(r"(?im)^\s*(?:formação(?: acadêmica)?|educação|escolaridade|education|academic background)\s*:?\s*$"),
    "skills": re.compile(r"(?im)^\s*(?:competências|habilidades|conhecimentos(?: técnicos)?|skills|technical skills|tecnologias)\s*:?\s*$"),
}
SPACED_LETTERS = re.compile(r"(?<!\w)(?:[A-Za-zÀ-ÿ]\s+){5,}[A-Za-zÀ-ÿ](?!\w)")


class AtsFinding(BaseModel):
    code: str
    severity: Severity
    title: str
    explanation: str
    action: str


class AtsReport(BaseModel):
    document_type: Kind
    assessment: Literal["review_needed", "basic_readability_confirmed"]
    readable_characters: int = Field(ge=0)
    findings: list[AtsFinding]
    limitations: list[str] = Field(default_factory=lambda: [
        "Esta checagem não simula o algoritmo privado de nenhum ATS e não garante seleção.",
        "A leitura bem-sucedida neste projeto não comprova leitura idêntica em outros sistemas.",
    ])


@dataclass(frozen=True)
class StructuralSignals:
    has_tables: bool = False
    has_images: bool = False
    has_header_or_footer_text: bool = False
    has_text_boxes: bool = False
    has_multiple_columns: bool = False


def _finding(
    code: str, severity: Severity, title: str, explanation: str, action: str
) -> AtsFinding:
    return AtsFinding(
        code=code, severity=severity, title=title, explanation=explanation, action=action
    )


def _docx_structure(payload: bytes) -> StructuralSignals:
    doc = Document(BytesIO(payload))
    from docx.oxml.ns import qn

    # XML-level checks, not assumptions based on document appearance.
    box = doc.element.body.xpath(".//w:txbxContent")
    cols = doc.element.body.xpath(".//w:cols")
    multiple_columns = any(
        int(value) > 1
        for node in cols
        if (value := node.get(qn("w:num"), "1")).isdigit()
    )
    header_text = any(
        any(p.text.strip() for p in section.header.paragraphs)
        or any(p.text.strip() for p in section.footer.paragraphs)
        for section in doc.sections
    )
    return StructuralSignals(
        has_tables=bool(doc.tables),
        has_images=bool(doc.inline_shapes),
        has_header_or_footer_text=header_text,
        has_text_boxes=bool(box),
        has_multiple_columns=multiple_columns,
    )


def _pdf_structure(payload: bytes) -> StructuralSignals:
    # PDF image objects are only a *signal*: decorative logos should not be
    # treated as proof that the whole document is image-only.
    reader = PdfReader(BytesIO(payload), strict=True)
    with_images = False
    for page in reader.pages:
        resources = page.get("/Resources")
        if resources is None:
            continue
        try:
            obj = resources.get_object()
            xobjects = obj.get("/XObject")
            if xobjects:
                images = xobjects.get_object()
                if any(
                    item.get_object().get("/Subtype") == "/Image"
                    for item in images.values()
                ):
                    with_images = True
                    break
        except (AttributeError, KeyError, TypeError, ValueError):
            continue
    return StructuralSignals(has_images=with_images)


def _report(text: str, kind: Kind, signals: StructuralSignals) -> AtsReport:
    findings: list[AtsFinding] = []
    if len(text.strip()) < 50:
        findings.append(_finding(
            "insufficient_text", "important", "Texto insuficiente para leitura",
            "Não foi possível recuperar conteúdo textual suficiente do documento.",
            "Use DOCX ou PDF com texto selecionável; documentos digitalizados podem exigir OCR.",
        ))
    if signals.has_multiple_columns:
        findings.append(_finding(
            "columns", "attention", "Layout com múltiplas colunas",
            "A ordem de leitura pode variar entre mecanismos de extração.",
            "Prefira uma única coluna para a versão enviada em candidaturas.",
        ))
    if signals.has_tables:
        findings.append(_finding(
            "tables", "attention", "Tabelas detectadas",
            "Alguns leitores podem reorganizar o texto contido em tabelas.",
            "Use seções e parágrafos simples em vez de tabelas para experiências e formação.",
        ))
    if signals.has_text_boxes:
        findings.append(_finding(
            "text_boxes", "attention", "Caixas de texto detectadas",
            "Caixas flutuantes podem ser ignoradas ou lidas fora da ordem.",
            "Coloque informações importantes diretamente no corpo do documento.",
        ))
    if signals.has_header_or_footer_text:
        findings.append(_finding(
            "headers", "attention", "Texto em cabeçalho ou rodapé",
            "Dados de contato nessas áreas podem não ser identificados pelo leitor.",
            "Mantenha nome e contato no corpo principal do currículo.",
        ))
    if signals.has_images:
        findings.append(_finding(
            "images", "suggestion", "Elementos gráficos presentes",
            "Imagens e ícones podem não ser interpretados como texto.",
            "Não coloque experiência ou contatos exclusivamente em elementos gráficos.",
        ))
    if text.strip() and not EMAIL_RE.search(text):
        findings.append(_finding(
            "email_missing", "attention", "E-mail não identificado",
            "Não encontramos um endereço de e-mail no texto extraído.",
            "Confira se o e-mail aparece como texto selecionável no corpo do currículo.",
        ))
    if text.strip() and not PHONE_RE.search(text):
        findings.append(_finding(
            "phone_missing", "suggestion", "Telefone não identificado",
            "Não encontramos um telefone em formato reconhecível no texto.",
            "Se fizer sentido para sua candidatura, informe um contato acessível no corpo.",
        ))
    if text.strip() and SPACED_LETTERS.search(text):
        findings.append(_finding(
            "spaced_letters", "attention", "Palavras com letras separadas",
            "Letras espaçadas artificialmente podem ser interpretadas como palavras diferentes.",
            "Substitua títulos com letras separadas por texto normal.",
        ))
    if len(text.strip()) >= 50:
        for code, pattern in SECTION_PATTERNS.items():
            if not pattern.search(text):
                titles = {
                    "experience": "Experiência profissional",
                    "education": "Formação",
                    "skills": "Competências",
                }
                findings.append(_finding(
                    "section_" + code, "suggestion", "Seção não identificada: " + titles[code],
                    "Não identificamos um título de seção convencional no texto extraído.",
                    "Se essa informação se aplica, use um título simples e explícito.",
                ))
    return AtsReport(
        document_type=kind,
        assessment="review_needed" if findings else "basic_readability_confirmed",
        readable_characters=len(text),
        findings=findings,
    )


def check_text(text: str) -> AtsReport:
    return _report(text, "text", StructuralSignals())


def check_document(filename: str | None, payload: bytes) -> AtsReport:
    kind = Path(filename or "").suffix.lower().lstrip(".")
    if kind not in {"pdf", "docx", "txt"}:
        raise ValueError("Formato inválido. Envie PDF, DOCX ou TXT.")
    if not payload or len(payload) > MAX_UPLOAD_BYTES:
        raise ValueError("O arquivo deve ter no máximo 3 MB.")
    # Run existing archive-size/page-count validation before inspecting
    # potentially expensive DOCX or PDF document structures.
    try:
        text = extract_resume(filename, payload)
    except PdfReadError as exc:
        raise ValueError("O arquivo PDF não pôde ser interpretado.") from exc
    except ValueError as exc:
        if "texto suficiente" not in str(exc):
            raise
        text = ""

    try:
        if kind == "docx":
            signals = _docx_structure(payload)
        elif kind == "pdf":
            signals = _pdf_structure(payload)
        else:
            signals = StructuralSignals()
    except (PdfReadError, ValueError, KeyError, TypeError, AttributeError) as exc:
        raise ValueError("Não foi possível verificar a estrutura do documento.") from exc
    return _report(text, kind, signals)
