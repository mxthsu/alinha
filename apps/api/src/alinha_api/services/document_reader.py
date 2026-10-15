"""In-memory extraction for user-provided files; nothing is persisted."""
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 3 * 1024 * 1024
MAX_EXTRACTED_CHARS = 22000
MAX_DOCX_UNCOMPRESSED_BYTES = 12 * 1024 * 1024
MAX_PDF_PAGES = 12
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_resume(filename: str | None, content: bytes) -> str:
    extension = Path(filename or "").suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Formato inválido. Envie PDF, DOCX ou TXT.")
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("O arquivo deve ter no máximo 3 MB.")

    try:
        if extension == ".pdf":
            if not content.startswith(b"%PDF"):
                raise ValueError("O arquivo não é um PDF válido.")
            reader = PdfReader(BytesIO(content), strict=True)
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ValueError("O PDF deve ter no máximo 12 páginas.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            if not content.startswith(b"PK"):
                raise ValueError("O arquivo não é um DOCX válido.")
            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise ValueError("O DOCX possui conteúdo descompactado excessivo.")
            document = Document(BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
        else:
            text = content.decode("utf-8-sig")
    except (BadZipFile, UnicodeDecodeError, OSError, KeyError) as exc:
        raise ValueError("Não foi possível ler o arquivo enviado.") from exc

    text = text.strip()
    if len(text) < 50:
        raise ValueError(
            "Não foi possível extrair texto suficiente. PDFs digitalizados ainda não são suportados."
        )
    if len(text) > MAX_EXTRACTED_CHARS:
        raise ValueError("O currículo excedeu o limite de texto permitido.")
    return text
