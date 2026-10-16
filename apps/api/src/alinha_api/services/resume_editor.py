"""Minimal, single-column DOCX generation for the locally approved resume text.

This function does not rewrite, classify or embellish qualifications. It
formats the exact non-empty lines entered by the candidate in a parseable
Word document and returns the bytes without persisting them.
"""
import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from alinha_api.services.document_reader import MAX_EXTRACTED_CHARS

HEADING_RE = re.compile(
    r"^(?:resumo(?: profissional)?|objetivo(?: profissional)?|perfil(?: profissional)?|"
    r"experi[eê]ncias? profissionais?|forma[cç][aã]o(?: acad[eê]mica)?|"
    r"educa[cç][aã]o|compet[eê]ncias|habilidades|conhecimentos(?: t[eé]cnicos)?|"
    r"projetos(?: relevantes)?|certifica[cç][oõ]es|idiomas|"
    r"professional summary|experience|work experience|education|"
    r"skills|technical skills|projects|certifications|languages)\s*:?$",
    re.IGNORECASE,
)


def export_resume_docx(resume_text: str) -> bytes:
    """Render manually edited content as selectable text in a table-free DOCX."""
    if not 50 <= len(resume_text.strip()) <= MAX_EXTRACTED_CHARS:
        raise ValueError("O currículo deve conter entre 50 e 22.000 caracteres de texto.")

    lines = [line.strip() for line in resume_text.splitlines() if line.strip()]
    if not lines:
        raise ValueError("O currículo não contém texto utilizável.")

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(32, 50, 40)
    style.paragraph_format.space_after = Pt(4)

    title_style = document.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(17)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(28, 62, 42)
    title_style.paragraph_format.space_after = Pt(5)

    heading_style = document.styles["Heading 1"]
    heading_style.font.name = "Arial"
    heading_style.font.size = Pt(11.5)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(31, 65, 46)
    heading_style.paragraph_format.space_before = Pt(11)
    heading_style.paragraph_format.space_after = Pt(5)
    heading_style.paragraph_format.keep_with_next = True

    for index, line in enumerate(lines):
        if index == 0:
            paragraph = document.add_paragraph(style="Title")
        elif HEADING_RE.fullmatch(line):
            paragraph = document.add_paragraph(style="Heading 1")
        else:
            paragraph = document.add_paragraph(style="Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.add_run(line)

    document.core_properties.title = "Currículo"
    document.core_properties.author = "Alinha"
    output = BytesIO()
    document.save(output)
    return output.getvalue()
