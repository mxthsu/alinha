"""Deterministic structured profile -> human-readable text and ATS-friendly DOCX.

The export repeats only candidate-supplied fields, never generates new claims.
"""
from io import BytesIO

from docx import Document
from docx.shared import Cm, Pt, RGBColor

from alinha_api.resume_schema import ResumeProfile, TemplateName


def render_profile_text(profile: ResumeProfile) -> str:
    sections: list[str] = []
    identity = [profile.contact.name.strip(), profile.contact.headline.strip()]
    sections.extend(line for line in identity if line)
    contacts = [
        profile.contact.email.strip(),
        profile.contact.phone.strip(),
        profile.contact.city.strip(),
        *[link.strip() for link in profile.contact.links],
    ]
    if any(contacts):
        sections.append(" | ".join(value for value in contacts if value))
    if profile.summary.strip():
        sections.extend(["Resumo profissional", profile.summary.strip()])
    if profile.experiences:
        sections.append("Experiência profissional")
        for item in profile.experiences:
            head = " — ".join(x.strip() for x in (item.role, item.organization) if x.strip())
            if head:
                sections.append(head)
            if item.period.strip():
                sections.append(item.period.strip())
            sections.extend("- " + x.strip() for x in item.highlights if x.strip())
    if profile.education:
        sections.append("Formação")
        for item in profile.education:
            line = " — ".join(x.strip() for x in (item.course, item.institution) if x.strip())
            if line:
                sections.append(line)
            if item.period.strip():
                sections.append(item.period.strip())
    if profile.projects:
        sections.append("Projetos")
        for item in profile.projects:
            sections.extend(
                x.strip() for x in (item.title, item.description, item.link) if x.strip()
            )
    if profile.skills:
        sections.extend(["Competências", ", ".join(x.strip() for x in profile.skills if x.strip())])
    if profile.languages:
        sections.extend(["Idiomas", ", ".join(x.strip() for x in profile.languages if x.strip())])
    return "\n".join(sections).strip()


def export_structured_docx(profile: ResumeProfile, template: TemplateName) -> bytes:
    text = render_profile_text(profile)
    if len(text) < 50 or len(text) > 22000:
        raise ValueError("Preencha informações suficientes, até 22.000 caracteres, para exportar.")
    document = Document()
    section = document.sections[0]
    section.left_margin = section.right_margin = Cm(2.0 if template == "classico" else 1.65)
    section.top_margin = section.bottom_margin = Cm(1.9 if template == "classico" else 1.45)
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5 if template == "classico" else 10)
    style.font.color.rgb = RGBColor(38, 49, 43)
    style.paragraph_format.space_after = Pt(5 if template == "classico" else 3)
    title = document.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(18 if template == "classico" else 16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(30, 61, 41)
    title.paragraph_format.space_after = Pt(5)
    heading = document.styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.size = Pt(12 if template == "classico" else 11)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(30, 61, 41)
    heading.paragraph_format.space_before = Pt(11 if template == "classico" else 7)
    heading.paragraph_format.space_after = Pt(4)

    headings = {
        "Resumo profissional", "Experiência profissional", "Formação",
        "Projetos", "Competências", "Idiomas",
    }
    for index, line in enumerate(text.splitlines()):
        style_name = "Title" if index == 0 else "Heading 1" if line in headings else "Normal"
        paragraph = document.add_paragraph(style=style_name)
        paragraph.add_run(line)
        if style_name == "Heading 1":
            paragraph.paragraph_format.keep_with_next = True

    document.core_properties.title = "Currículo"
    document.core_properties.author = "Alinha"
    output = BytesIO()
    document.save(output)
    return output.getvalue()
