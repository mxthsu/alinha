"""ATS checks are synthetic and require no Gemini key or provider requests."""
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from alinha_api.main import app
from alinha_api.services.ats import check_document, check_text


def _sample_resume() -> str:
    return (
        "Pessoa Desenvolvedora\n"
        "Contato: pessoa@exemplo.com | +55 48 99999-1234\n"
        "Experiência profissional\n"
        "Trabalhou em projetos fictícios com Python e FastAPI.\n"
        "Formação\n"
        "Curso técnico fictício de desenvolvimento.\n"
        "Competências\n"
        "Python, documentação e testes automatizados."
    )


def test_plain_text_has_no_universal_ats_score() -> None:
    report = check_text(_sample_resume())
    assert report.document_type == "text"
    assert report.assessment == "basic_readability_confirmed"
    assert report.readable_characters > 50
    assert not report.findings
    assert "score" not in report.model_dump()
    assert any("não garante" in notice for notice in report.limitations)


def test_missing_sections_and_contact_are_actionable() -> None:
    report = check_text(
        "Resumo de projetos fictícios de desenvolvimento Python e construção "
        "de sistemas documentados para aprendizado profissional."
    )
    codes = {finding.code for finding in report.findings}
    assert {"email_missing", "section_experience", "section_education"} <= codes
    assert report.assessment == "review_needed"


def test_docx_table_and_header_are_detected() -> None:
    document = Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "Pessoa Fictícia | pessoa@exemplo.com"
    for line in _sample_resume().splitlines():
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Projeto fictício"
    table.cell(0, 1).text = "Python"
    buffer = BytesIO()
    document.save(buffer)
    report = check_document("curriculo.docx", buffer.getvalue())
    codes = {finding.code for finding in report.findings}
    assert "tables" in codes
    assert "headers" in codes
    assert report.document_type == "docx"


def test_image_only_pdf_warns_without_ocr() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    output = BytesIO()
    writer.write(output)
    report = check_document("curriculo.pdf", output.getvalue())
    assert report.document_type == "pdf"
    assert report.readable_characters == 0
    assert "insufficient_text" in {item.code for item in report.findings}


def test_file_size_and_extension_checked() -> None:
    for filename, payload in [
        ("bad.exe", b"binary!"),
        ("resume.pdf", b""),
        ("resume.txt", b"x" * (3 * 1024 * 1024 + 1)),
    ]:
        try:
            check_document(filename, payload)
        except ValueError:
            continue
        assert False, f"Accepted invalid upload: {filename}"


def test_api_works_without_an_ai_connection() -> None:
    client = TestClient(app)
    response = client.post("/api/v1/ats/check/text", json={"resume_text": _sample_resume()})
    assert response.status_code == 200
    assert response.json()["assessment"] == "basic_readability_confirmed"
    assert response.json()["findings"] == []


def test_api_reports_scanned_pdf_instead_of_500() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    output = BytesIO()
    writer.write(output)
    client = TestClient(app)
    response = client.post(
        "/api/v1/ats/check/upload",
        files={"file": ("cv.pdf", output.getvalue(), "application/pdf")},
    )
    assert response.status_code == 200
    assert response.json()["findings"][0]["code"] == "insufficient_text"
