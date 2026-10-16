"""Round-trip tests for document editing/export with fictional CVs only."""
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from alinha_api.main import app
from alinha_api.services.ats import check_document
from alinha_api.services.document_reader import extract_resume
from alinha_api.services.resume_editor import export_resume_docx

RESUME = (
    "Pessoa Desenvolvedora Fictícia\n"
    "pessoa@exemplo.com | (48) 99999-1234\n"
    "Experiência profissional\n"
    "Desenvolveu APIs Python e FastAPI em um projeto fictício.\n"
    "Formação\n"
    "Graduação fictícia em tecnologia.\n"
    "Competências\n"
    "Python, documentação de sistemas, testes com pytest.\n"
)


def test_export_is_single_column_textual_and_complete():
    data = export_resume_docx(RESUME)
    document = Document(BytesIO(data))
    original_lines = [item.strip() for item in RESUME.splitlines() if item.strip()]
    exported_lines = [p.text for p in document.paragraphs if p.text.strip()]
    assert original_lines == exported_lines
    assert not document.tables
    assert not document.inline_shapes
    assert document.paragraphs[0].style.name == "Title"
    assert "Heading 1" in [p.style.name for p in document.paragraphs]
    assert not document.sections[0].header.paragraphs[0].text
    assert not document.sections[0].footer.paragraphs[0].text
    assert extract_resume("resume.docx", data).splitlines() == original_lines
    assert check_document("resume.docx", data).readable_characters > 50


def test_api_import_and_export_roundtrip():
    client = TestClient(app)
    original = export_resume_docx(RESUME)
    imported = client.post(
        "/api/v1/ats/editor/import",
        files={"file": ("ficticio.docx", original, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert imported.status_code == 200
    assert imported.json()["resume_text"].splitlines() == RESUME.strip().splitlines()
    exported = client.post(
        "/api/v1/ats/editor/export/docx",
        json={"resume_text": imported.json()["resume_text"]},
    )
    assert exported.status_code == 200
    assert exported.content.startswith(b"PK")
    assert exported.headers["cache-control"] == "no-store"
    assert "attachment" in exported.headers["content-disposition"]
    assert "pessoa@exemplo.com" in extract_resume("cv.docx", exported.content)


def test_no_synthetic_qualifications_or_templates_added():
    text = (
        "Pessoa Desenvolvedora\n"
        "pessoa@exemplo.com\n"
        "Experiência profissional\n"
        "Projeto pessoal fictício usando Python e testes.\n"
    )
    content = export_resume_docx(text)
    round_trip = extract_resume("result.docx", content)
    assert "Kubernetes" not in round_trip
    assert "SQL" not in round_trip
    assert round_trip.splitlines() == text.strip().splitlines()


def test_export_rejects_empty_document():
    response = TestClient(app).post(
        "/api/v1/ats/editor/export/docx",
        json={"resume_text": "x" * 20},
    )
    assert response.status_code == 422


def test_import_rejects_invalid_pdf():
    response = TestClient(app).post(
        "/api/v1/ats/editor/import",
        files={"file": ("ficticio.pdf", b"not actually PDF", "application/pdf")},
    )
    assert response.status_code == 422


def test_malformed_pdf_header_is_rejected_safely():
    response = TestClient(app).post(
        "/api/v1/ats/editor/import",
        files={"file": ("cv.pdf", b"%PDF-1.7\\nnot a valid document", "application/pdf")},
    )
    assert response.status_code == 422
    assert "PDF" in response.json()["detail"]
