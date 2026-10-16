"""Synthetic layout regression tests. Never include real candidate information."""
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from alinha_api.main import app
from alinha_api.services.structured_import import draft_profile

MODERN = """CANDIDATA FICTÍCIA
Engenheira de Software | Python Backend | IA aplicada
Cidade Exemplo, SC | +55 48 99999-8888 | candidata@exemplo.test | linkedin.com/in/candidata-
ficticia
RESUMO PROFISSIONAL
Experiência fictícia desenvolvendo soluções.
Trabalho com backend e automação.
EXPERIÊNCIA PROFISSIONAL
Engenheira de Software | Backend e Integrações | Empresa Exemplo do
Brasil Ltda.
ago/2024 - atual | Cidade Exemplo, SC
• Implementei integrações seguras usando Python,
FastAPI e PostgreSQL.
• Automatizei a verificação de contratos.
Analista de Sistemas | Empresa Fictícia
jan/2022 - jul/2024 | Brasil
• Desenvolvi APIs REST.
FORMAÇÃO
Bacharelado em Sistemas de Informação | Universidade Exemplo | concluído em jul/2024
Técnico em Informática | Instituto Fictício
COMPETÊNCIAS TÉCNICAS
Backend: Python, FastAPI, SQL, PostgreSQL.
Inteligência artificial: LLMs, RAG, embeddings,
Azure OpenAI, Gemini.
Sistemas: RabbitMQ, Redis, webhooks.
PROJETOS SELECIONADOS
Monitoramento de filas | Projeto acadêmico fictício
• Criação de pipeline de dados com Python.
• Inclusão de métricas e auditoria.
Plataforma de recrutamento | Projeto interno fictício
• API com FastAPI e banco relacional.
IDIOMAS
Português: nativo | Inglês: avançado | Espanhol: básico
"""


def test_modern_wrapped_layout_populates_every_studio_section():
    result = draft_profile(MODERN)
    profile = result.profile
    assert profile.contact.name == "CANDIDATA FICTÍCIA"
    assert profile.contact.headline == "Engenheira de Software | Python Backend | IA aplicada"
    assert profile.contact.email == "candidata@exemplo.test"
    assert profile.contact.phone == "+55 48 99999-8888"
    assert profile.contact.city == "Cidade Exemplo, SC"
    assert profile.contact.links == ["linkedin.com/in/candidata-ficticia"]
    assert "Trabalho com backend" in profile.summary
    assert len(profile.experiences) == 2
    assert profile.experiences[0].organization == "Empresa Exemplo do Brasil Ltda."
    assert profile.experiences[0].period == "ago/2024 - atual"
    assert profile.experiences[0].highlights[0].endswith("FastAPI e PostgreSQL.")
    assert len(profile.education) == 2
    assert profile.education[0].period == "concluído em jul/2024"
    assert profile.education[1].institution == "Instituto Fictício"
    assert len(profile.projects) == 2
    assert "auditoria" in profile.projects[0].description
    assert "Python" in profile.skills and "Azure OpenAI" in profile.skills
    assert "Backend" not in profile.skills
    assert profile.languages == ["Português: nativo", "Inglês: avançado", "Espanhol: básico"]
    assert result.unassigned_lines == []


def test_classic_dash_layout_and_unknown_section_are_kept_safe():
    text = """Candidate Example
candidate@example.test
Professional Summary
Software developer with fictional experience.
Work Experience
Developer — Fictional Company — 2022–2025
- Built internal tools.
Education
Computer Science — Example University — 2018–2022
Projects
Portfolio: Open source sample project.
Technical Skills
Python, Docker, Git
Certifications
Unverified certificate details
"""
    result = draft_profile(text)
    assert result.profile.experiences[0].organization == "Fictional Company"
    assert result.profile.education[0].institution == "Example University"
    assert result.profile.projects[0].title == "Portfolio"
    assert result.profile.skills == ["Python", "Docker", "Git"]
    assert "Unverified certificate details" in result.unassigned_lines


def test_unclassifiable_roles_stay_unassigned_without_invented_experience():
    text = """Pessoa de Teste
pessoa@exemplo.test
Experiência profissional
Desempenhei funções variadas em vários contextos.
Competências
Python, SQL
"""
    result = draft_profile(text)
    assert not result.profile.experiences
    assert "Desempenhei funções variadas em vários contextos." in result.unassigned_lines


def test_import_docx_with_wrapped_layout_populates_sections():
    doc = Document()
    for line in MODERN.splitlines():
        doc.add_paragraph(line)
    out = BytesIO()
    doc.save(out)
    response = TestClient(app).post(
        "/api/v1/resumes/structured/import/upload",
        files={"file": ("ficticio.docx", out.getvalue(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    data = response.json()
    assert len(data["profile"]["experiences"]) == 2
    assert len(data["profile"]["projects"]) == 2
    assert "FastAPI" in data["profile"]["skills"]


def _minimal_pdf(lines: list[str]) -> bytes:
    """Generate an ASCII-only PDF fixture with no third-party PDF creation library."""
    commands = ["BT /F1 10 Tf 50 770 Td 13 TL"]
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(f"({escaped}) Tj T*")
    commands.append("ET")
    stream = "\n".join(commands).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Root 1 0 R /Size {len(offsets)} >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(output)


def test_pdf_import_pipeline_with_headings_and_fields():
    sample = [
        "Candidate Example", "Backend Engineer | Python", "city@example.test",
        "SUMMARY", "Developed sample APIs and backend tools.",
        "WORK EXPERIENCE", "Backend Developer | Example Company | 2022 - 2025",
        "- Built APIs and automated operations.",
        "EDUCATION", "Computer Science | Sample College | 2018 - 2022",
        "SELECTED PROJECTS", "API Portfolio: Sample REST backend project.",
        "TECHNICAL SKILLS", "Python, FastAPI, SQL",
        "LANGUAGES", "English, Spanish",
    ]
    response = TestClient(app).post(
        "/api/v1/resumes/structured/import/upload",
        files={"file": ("fictional.pdf", _minimal_pdf(sample), "application/pdf")},
    )
    assert response.status_code == 200, response.text
    profile = response.json()["profile"]
    assert profile["contact"]["name"] == "Candidate Example"
    assert profile["experiences"][0]["organization"] == "Example Company"
    assert profile["education"][0]["institution"] == "Sample College"
    assert profile["projects"][0]["title"] == "API Portfolio"
    assert profile["skills"] == ["Python", "FastAPI", "SQL"]


def test_docx_table_preserves_section_order():
    doc = Document()
    doc.add_paragraph("Pessoa Fictícia")
    doc.add_paragraph("Experiência profissional")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Analista Backend"
    table.cell(0, 1).text = "Empresa Exemplo"
    doc.add_paragraph("Competências")
    doc.add_paragraph("Python, SQL")
    out = BytesIO()
    doc.save(out)
    response = TestClient(app).post(
        "/api/v1/resumes/structured/import/upload",
        files={"file": ("table.docx", out.getvalue(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200, response.text
    profile = response.json()["profile"]
    assert profile["experiences"][0]["role"] == "Analista Backend"
    assert profile["experiences"][0]["organization"] == "Empresa Exemplo"
    assert profile["skills"] == ["Python", "SQL"]


def test_txt_windows_encoding_with_accents_is_readable():
    text = "Pessoa Fictícia\npessoa@exemplo.test\nCompetências técnicas\nPython, Automação\n"
    response = TestClient(app).post(
        "/api/v1/resumes/structured/import/upload",
        files={"file": ("resume.txt", text.encode("cp1252"), "text/plain")},
    )
    assert response.status_code == 200
    assert response.json()["profile"]["skills"] == ["Python", "Automação"]


def test_scanned_or_empty_pdf_returns_actionable_error():
    response = TestClient(app).post(
        "/api/v1/resumes/structured/import/upload",
        files={"file": ("empty.pdf", _minimal_pdf([]), "application/pdf")},
    )
    assert response.status_code == 422
    assert "PDFs digitalizados" in response.json()["detail"]
