"""Structured resume & grounded review tests, exclusively synthetic data."""
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from alinha_api.main import app
from alinha_api.resume_schema import ResumeProfile
from alinha_api.schemas import AnalysisResult, RequirementDraft
from alinha_api.services.structured_resume import export_structured_docx, render_profile_text
from alinha_api.workflows.resume_review import build_resume_review_workflow

EXAMPLE = {
    "schema_version": 1,
    "contact": {
        "name": "Pessoa Exemplo",
        "headline": "Desenvolvedora Python",
        "email": "exemplo@dominio.test",
        "city": "Cidade fictícia",
    },
    "summary": "Desenvolve APIs REST com Python e FastAPI em projetos acadêmicos.",
    "experiences": [
        {
            "role": "Pessoa Desenvolvedora",
            "organization": "Organização fictícia",
            "period": "2025 - 2026",
            "highlights": ["Desenvolveu APIs com Python e FastAPI.", "Documentou endpoints REST."],
        }
    ],
    "education": [{"course": "Tecnologia", "institution": "Universidade fictícia"}],
    "projects": [{"title": "Projeto de integração", "description": "Serviços REST em Python."}],
    "skills": ["Python", "FastAPI", "SQL"],
    "languages": ["Inglês intermediário"],
}


def test_schema_and_text_export_preserve_only_submitted_values():
    profile = ResumeProfile.model_validate(EXAMPLE)
    text = render_profile_text(profile)
    for fragment in ["Pessoa Exemplo", "Python", "2025 - 2026", "Projeto de integração"]:
        assert fragment in text
    assert "Kubernetes" not in text
    for template in ("classico", "compacto"):
        raw = export_structured_docx(profile, template)
        document = Document(BytesIO(raw))
        extracted = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert extracted == text
        assert not document.tables
        assert len(document.inline_shapes) == 0
        assert "Heading 1" in [item.style.name for item in document.paragraphs]
        assert "exemplo@dominio.test" in extracted


def test_api_preview_and_export_without_ai():
    client = TestClient(app)
    preview = client.post("/api/v1/resumes/structured/preview/text", json=EXAMPLE)
    assert preview.status_code == 200
    assert "Pessoa Exemplo" in preview.json()["resume_text"]
    export = client.post(
        "/api/v1/resumes/structured/export/docx",
        json={"profile": EXAMPLE, "template": "compacto"},
    )
    assert export.status_code == 200
    assert export.content.startswith(b"PK")
    assert export.headers["cache-control"] == "no-store"


def test_empty_profile_not_exported():
    response = TestClient(app).post(
        "/api/v1/resumes/structured/export/docx", json={"profile": {}}
    )
    assert response.status_code == 422


def test_invalid_profile_does_not_persist_or_invent_data():
    response = TestClient(app).post(
        "/api/v1/resumes/structured/export/docx",
        json={"profile": {**EXAMPLE, "skills": ["invented " * 60]}},
    )
    assert response.status_code == 422


def test_review_agent_one_ai_call_and_human_approval():
    profile = ResumeProfile.model_validate(EXAMPLE)
    resume = render_profile_text(profile)
    job = "Vaga fictícia exige experiência com Python e FastAPI e Kubernetes em produção."
    calls = []

    def analyzer(resume_text, job_description):
        calls.append((resume_text, job_description))
        return AnalysisResult(
            job_title="Backend",
            summary="Resumo das evidências verificadas.",
            requirements=[
                RequirementDraft(
                    requirement="Python e FastAPI",
                    job_evidence="Python e FastAPI",
                    resume_evidence="Python e FastAPI",
                    status="evidenciado",
                    recommendation="Destaque os projetos reais com Python e FastAPI.",
                ),
                RequirementDraft(
                    requirement="Kubernetes",
                    job_evidence="Kubernetes em produção",
                    resume_evidence=None,
                    status="nao_identificado",
                    recommendation="Não declare Kubernetes sem experiência comprovada.",
                ),
            ],
            limitations=["Ferramenta orientativa."],
        )

    output = build_resume_review_workflow(analyzer).invoke({
        "resume_text": resume, "job_description": job,
    })["review"]
    assert len(calls) == 1
    assert len(output.suggestions) == 2
    assert all(item.approval_required for item in output.suggestions)
    assert profile.model_dump() == ResumeProfile.model_validate(EXAMPLE).model_dump()


def test_review_agent_filters_invented_citations():
    def suspicious(*_):
        return AnalysisResult(
            job_title="Backend", summary="Relatório de teste",
            requirements=[
                RequirementDraft(
                    requirement="AWS",
                    job_evidence="AWS avançado",
                    resume_evidence="Certificação falsa",
                    status="evidenciado",
                    recommendation="Adicione AWS.",
                ),
            ], limitations=[],
        )

    result = build_resume_review_workflow(suspicious).invoke({
        "resume_text": "Trabalhei em Python",
        "job_description": "Vaga requer Python",
    })["review"]
    assert result.suggestions == []


def test_review_route_checks_profile_before_provider_call(monkeypatch):
    monkeypatch.setattr("alinha_api.main.run_analysis", lambda _: (_ for _ in ()).throw(
        AssertionError("The model should not have been called")
    ))
    client = TestClient(app)
    response = client.post(
        "/api/v1/resumes/structured/review",
        json={"profile": {}, "job_description": "Vaga fictícia procura Python e FastAPI, testes e documentação profissional em APIs."},
    )
    assert response.status_code == 422
